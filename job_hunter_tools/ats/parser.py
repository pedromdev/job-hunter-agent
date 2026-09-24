import os
from pathlib import Path

_IMPORT_ERRORS = {}


def extract_text(file_path: str) -> str:
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".pdf":
        return _extract_pdf(file_path)
    elif ext == ".txt":
        return _extract_txt(file_path)
    elif ext in (".docx", ".doc"):
        return _extract_docx(file_path)
    else:
        raise ValueError(f"Formato não suportado: '{ext}'. Use .pdf, .txt ou .docx")


def _extract_pdf(pdf_path: str) -> str:
    try:
        from PyPDF2 import PdfReader
    except ImportError:
        raise ImportError("PyPDF2 é necessário para extrair PDFs. Instale com: pip install PyPDF2")
    text = ""
    reader = PdfReader(pdf_path)
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + " "
    if not text.strip():
        raise ValueError(
            f"Não foi possível extrair texto de '{pdf_path}'. "
            "O PDF pode ser baseado em imagem (scaneado)."
        )
    return text


def _extract_txt(txt_path: str) -> str:
    return Path(txt_path).read_text(encoding="utf-8", errors="ignore")


def _extract_docx(docx_path: str) -> str:
    try:
        import docx
    except ImportError:
        raise ImportError("python-docx é necessário para extrair DOCX. Instale com: pip install python-docx")
    doc = docx.Document(docx_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)
    return " ".join(paragraphs)
